"""
表格动态扩展子模块 - 渲染引擎拆分
功能：根据数据行数动态增减表格行，保留表头样式，自动行高适配
依赖：python-pptx
"""
from copy import deepcopy
from typing import Any

from pptx.util import Inches

from aippt.logger import logger


def fill_dynamic_table(shape: Any, table_data: dict[str, Any]) -> None:
    """动态填充表格数据，保留第 0 行表头样式模板

    - 根据数据行数动态增减行（克隆最后一行保持样式一致性）
    - 填入表头数据（覆盖模板占位文本，保留样式）
    - 填入数据行，每行单元格继承样式（字体/对齐/背景色）
    - 自动行高适配

    :param shape: GraphicFrame 形状，含 table
    :param table_data: 表格数据，结构为
        {"headers": [...], "rows": [[...], ...]}
    """
    try:
        table = shape.table
    except Exception as e:
        logger.warning("无法访问表格数据: %s", e)
        return

    headers = table_data.get('headers', [])
    rows = table_data.get('rows', [])
    if not headers:
        logger.warning("表格 headers 为空，跳过填充")
        return

    current_rows = len(table.rows)
    if current_rows == 0:
        logger.warning("表格无行数据，无法填充")
        return
    needed_rows = len(rows) + 1  # +1 为表头行

    # 动态增减行（python-pptx 的 _RowCollection 不支持负索引，需用正索引）
    if needed_rows > current_rows:
        # 新增行：以最后一行为模板克隆（保留行结构与单元格样式）
        last_tr = table.rows[current_rows - 1]._tr
        added = 0
        for _ in range(needed_rows - current_rows):
            new_tr = deepcopy(last_tr)
            last_tr.addnext(new_tr)
            last_tr = new_tr
            added += 1
        logger.info("表格新增 %d 行（%d → %d）", added, current_rows, needed_rows)
    elif needed_rows < current_rows:
        # 删除多余行（从末尾删除，保留表头）
        extra = current_rows - needed_rows
        for _ in range(extra):
            remaining = len(table.rows)
            last_tr = table.rows[remaining - 1]._tr
            last_tr.getparent().remove(last_tr)
        logger.info("表格删除 %d 行（%d → %d）", extra, current_rows, needed_rows)

    # 填入表头数据（第 0 行，保留单元格样式）
    col_count = len(table.columns)
    for col_idx, header_text in enumerate(headers):
        if col_idx < col_count:
            cell = table.cell(0, col_idx)
            set_cell_text(cell, str(header_text))

    # 填入数据行
    for row_idx, row_data in enumerate(rows, start=1):
        if row_idx >= len(table.rows):
            break
        # 自动行高适配：默认 0.4 英寸，内容较长时增至 0.6
        try:
            max_cell_len = max((len(str(c)) for c in row_data), default=0)
            row_height = Inches(0.6) if max_cell_len > 20 else Inches(0.4)
            table.rows[row_idx].height = row_height
        except Exception:
            pass
        for col_idx, cell_text in enumerate(row_data):
            if col_idx < col_count:
                cell = table.cell(row_idx, col_idx)
                set_cell_text(cell, str(cell_text))

    logger.info("表格填充完成（表头=%d 列, 数据行=%d）",
                min(len(headers), col_count), len(rows))


def set_cell_text(cell: Any, text: str) -> None:
    """设置单元格文本，保留首个 run 的格式（字体/字号/颜色/粗体）

    :param cell: pptx table Cell 对象
    :param text: 要填入的文本
    """
    tf = cell.text_frame
    paragraphs = list(tf.paragraphs)
    if not paragraphs:
        cell.text = str(text)
        return
    first_para = paragraphs[0]
    if first_para.runs:
        # 保留首个 run 的格式，仅替换其 text，删除同段其余 run
        first_para.runs[0].text = str(text)
        for run in first_para.runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        # 没有现有 run，直接设置（会创建新 run，继承段落/单元格级格式）
        cell.text = str(text)
    # 删除多余段落（保留首段段落属性）
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)
